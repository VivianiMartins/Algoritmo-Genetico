from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _converter_cargo_numerico_para_texto(cargo_numerico):
    """
    Converte código numérico de cargo para texto.

    Entrada:
      cargo_numerico: int (1 = Farmacêutico, 2 = Auxiliar)

    Retorna:
      str: Nome do cargo em texto
    """
    mapa_cargos = {
        1: "Farmacêutico",
        2: "Auxiliar"
    }
    return mapa_cargos.get(cargo_numerico, "Desconhecido")


def salvar_excel( populacao, lista_nomes_cargos, farmaceuticos, auxiliares, caminho_arquivo='populacao_escalas.xlsx', num_dias_mes=30,
                  salvar_completo=True, indice_individuo=None, somatorio=False):
    """
    Salva em arquivo Excel com opção de salvar população completa ou apenas uma escala.
      Arquivo Excel com abas conforme a flag:
      - Se salvar_completo=True: múltiplas abas ("Indivíduo_1", "Indivíduo_2", ...)
      - Se salvar_completo=False: uma aba ("Escala")
    """

    qtdeFuncionarios, qtdeDias, qtdeIndividuos = populacao.shape

    # Extrai nomes e cargos da lista
    nomes_funcionarios = []
    cargos_texto = []
    ultimo_dia_mes_anterior = []

    for item in lista_nomes_cargos:
        if isinstance(item, tuple):
            nome, cargo_num, ultimo_dia_mes_ant = item
        elif isinstance(item, dict):
            nome = item.get('nome', 'Desconhecido')
            cargo_num = item.get('cargo', 2)
            ultimo_dia_mes_ant = item.get('ultimo_dia_mes_ant', 0)
        else:
            nome = str(item)
            cargo_num = 2
            ultimo_dia_mes_ant = 0

        nomes_funcionarios.append(nome)
        cargos_texto.append(_converter_cargo_numerico_para_texto(cargo_num))
        ultimo_dia_mes_anterior.append(ultimo_dia_mes_ant)

    # Valida quantidade de funcionários
    if len(nomes_funcionarios) != qtdeFuncionarios:
        raise ValueError(
            f"Número de funcionários na lista ({len(nomes_funcionarios)}) "
            f"não corresponde à população ({qtdeFuncionarios})"
        )

    # Define qual indivíduo salvar se for modo individual
    if not salvar_completo:
        if indice_individuo is None:
            indice_individuo = qtdeIndividuos - 1  # Último indivíduo
        elif indice_individuo >= qtdeIndividuos or indice_individuo < 0:
            raise ValueError(f"Índice inválido. Deve estar entre 0 e {qtdeIndividuos - 1}")

    # Cria workbook Excel
    workbook = Workbook()
    workbook.remove(workbook.active)  # Remove aba padrão

    # ===== DEFINIÇÃO DE ESTILOS =====
    cor_trabalha = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")  # Verde
    cor_nao_trabalha = PatternFill(start_color="FFB6C6", end_color="FFB6C6", fill_type="solid")  # Rosa
    cor_cabecalho = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")  # Azul
    cor_soma = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Amarelo

    fonte_negrito = Font(bold=True)
    fonte_cabecalho = Font(bold=True, color="FFFFFF")
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    borda = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # ===== FUNÇÃO AUXILIAR: Cria uma aba com a escala =====
    def criar_aba_escala(sheet, populacao_individuo, nomes, cargos, ultimo_dia_mes_anterior, farmaceuticos, auxiliares):
        """
        Cria uma aba formatada com a escala de um indivíduo.

        Entrada:
          sheet: worksheet do openpyxl
          populacao_individuo: array 2D [funcionários x dias]
          nomes: lista com nomes dos funcionários
          cargos: lista com cargos em texto dos funcionários
          farmaceuticos: número total de farmacêuticos
          auxiliares: número total de auxiliares
        """
        # ===== CABEÇALHO (LINHA 1) =====
        # Coluna A: Funcionário
        sheet['A1'] = 'Funcionário'
        sheet['A1'].font = fonte_cabecalho
        sheet['A1'].alignment = alinhamento_centro
        sheet['A1'].fill = cor_cabecalho

        # Coluna B: Cargo
        sheet['B1'] = 'Cargo'
        sheet['B1'].font = fonte_cabecalho
        sheet['B1'].alignment = alinhamento_centro
        sheet['B1'].fill = cor_cabecalho

        # Coluna C: Último dia do mês anterior
        sheet['C1'] = 'Último dia do mês anterior'
        sheet['C1'].font = fonte_cabecalho
        sheet['C1'].alignment = alinhamento_centro
        sheet['C1'].fill = cor_cabecalho

        # Colunas D em diante: Dias do mês
        for dia in range(1, qtdeDias + 1):
            coluna = get_column_letter(dia + 3)  # D, E, ...
            celula = sheet[f'{coluna}1']
            celula.value = f'Dia {dia}'
            celula.font = fonte_cabecalho
            celula.fill = cor_cabecalho
            celula.alignment = alinhamento_centro

        # Coluna total (se somatório ativo)
        if somatorio:
            coluna_total = get_column_letter(qtdeDias + 4)
            sheet[f'{coluna_total}1'] = 'Total dias escalados'
            sheet[f'{coluna_total}1'].font = fonte_cabecalho
            sheet[f'{coluna_total}1'].fill = cor_cabecalho
            sheet[f'{coluna_total}1'].alignment = alinhamento_centro

        # ===== DADOS DOS FUNCIONÁRIOS (LINHAS 2+) =====
        for func in range(qtdeFuncionarios):
            linha = func + 2

            # Coluna A: Nome
            sheet[f'A{linha}'] = nomes[func]
            sheet[f'A{linha}'].font = fonte_negrito
            sheet[f'A{linha}'].alignment = alinhamento_centro

            # Coluna B: Cargo (em texto)
            sheet[f'B{linha}'] = cargos[func]
            sheet[f'B{linha}'].alignment = alinhamento_centro
            sheet[f'B{linha}'].border = borda

            # Coluna C: Último dia do mês anterior (0 ou 1)
            sheet[f'C{linha}'] = ultimo_dia_mes_anterior[func]
            sheet[f'C{linha}'].alignment = alinhamento_centro
            sheet[f'C{linha}'].border = borda

            # Colunas de dias
            for dia in range(qtdeDias):
                coluna = get_column_letter(dia + 4)  # D, E, ...
                celula = sheet[f'{coluna}{linha}']

                valor = int(populacao_individuo[func, dia])
                celula.value = valor
                celula.alignment = alinhamento_centro
                celula.border = borda

                # Cor: Verde (1) ou Rosa (0)
                celula.fill = cor_trabalha if valor == 1 else cor_nao_trabalha

            # Fórmula: Total de dias escalados (se somatório ativo)
            if somatorio:
                coluna_total = get_column_letter(qtdeDias + 4)
                celula_total = sheet[f'{coluna_total}{linha}']
                primeira_coluna = get_column_letter(4)  # D
                ultima_coluna = get_column_letter(qtdeDias + 3)
                celula_total.value = f'=SUM({primeira_coluna}{linha}:{ultima_coluna}{linha})'
                celula_total.data_type = 'f'
                celula_total.alignment = alinhamento_centro
                celula_total.border = borda
                celula_total.font = fonte_negrito
                celula_total.fill = cor_soma

        # ===== AJUSTA LARGURA DAS COLUNAS =====
        sheet.column_dimensions['A'].width = 18  # Nome
        sheet.column_dimensions['B'].width = 14  # Cargo
        sheet.column_dimensions['C'].width = 24  # Último dia do mês anterior

        for dia in range(1, qtdeDias + 2):
            coluna = get_column_letter(dia + 3)
            sheet.column_dimensions[coluna].width = 8

        if somatorio:
            coluna_total = get_column_letter(qtdeDias + 4)
            sheet.column_dimensions[coluna_total].width = 16

        # ===== RODAPÉ COM SOMATÓRIOS (Se ativo) =====
        if somatorio:
            # Identifica as linhas exatas onde vão ficar os somatórios no Excel
            linha_rodape_farm = qtdeFuncionarios + 2
            linha_rodape_aux = qtdeFuncionarios + 3

            # Rótulos das linhas de somatório
            sheet[f'A{linha_rodape_farm}'] = 'Total de farmacêuticos:'
            sheet[f'A{linha_rodape_farm}'].font = fonte_negrito
            sheet[f'A{linha_rodape_farm}'].fill = cor_soma

            sheet[f'A{linha_rodape_aux}'] = 'Total de auxiliares:'
            sheet[f'A{linha_rodape_aux}'].font = fonte_negrito
            sheet[f'A{linha_rodape_aux}'].fill = cor_soma

            # Define o intervalo de linhas (range) de cada grupo para a fórmula do Excel
            inicio_farm = 2
            fim_farm = 1 + farmaceuticos

            inicio_aux = fim_farm + 1
            fim_aux = fim_farm + auxiliares

            for dia in range(qtdeDias):
                coluna = get_column_letter(dia + 4)  # Colunas D, E, etc.

                # Somatório Farmacêuticos
                celula_farm = sheet[f'{coluna}{linha_rodape_farm}']
                celula_farm.value = f'=SUM({coluna}{inicio_farm}:{coluna}{fim_farm})'
                celula_farm.data_type = 'f'
                celula_farm.alignment = alinhamento_centro
                celula_farm.font = Font(bold=True)
                celula_farm.fill = cor_soma
                celula_farm.border = borda

                # Somatório Auxiliares
                celula_aux = sheet[f'{coluna}{linha_rodape_aux}']
                celula_aux.value = f'=SUM({coluna}{inicio_aux}:{coluna}{fim_aux})'
                celula_aux.data_type = 'f'
                celula_aux.alignment = alinhamento_centro
                celula_aux.font = Font(bold=True)
                celula_aux.fill = cor_soma
                celula_aux.border = borda

    # ===== LÓGICA PRINCIPAL =====
    if salvar_completo:
        for ind in range(qtdeIndividuos):
            sheet = workbook.create_sheet(title=f'Indivíduo_{ind + 1}')
            criar_aba_escala(sheet, populacao[:, :, ind], nomes_funcionarios, cargos_texto, ultimo_dia_mes_anterior, farmaceuticos, auxiliares)
        workbook.save(caminho_arquivo)

    else: # Salva apenas um indivíduo
        sheet = workbook.create_sheet(title='Escala')
        criar_aba_escala(sheet, populacao[:, :, indice_individuo], nomes_funcionarios, cargos_texto, ultimo_dia_mes_anterior, farmaceuticos, auxiliares)
        workbook.save(caminho_arquivo)